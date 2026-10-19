"""Тесты вывода: форматирование, печать и структура Excel-файла."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from clean import CANONICAL_COLUMNS, prepare
from output import Report, format_money, print_summary, save_excel
from report import summarize

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_data"


@pytest.fixture(scope="module")
def report() -> Report:
    df, stats = prepare(
        [
            SAMPLE_DIR / "sales_january.csv",
            SAMPLE_DIR / "sales_february.csv",
            SAMPLE_DIR / "sales_march.csv",
        ]
    )
    summary = summarize(df, "Товар")
    return Report(
        summary=summary,
        data=df,
        stats=stats,
        source_names=["sales_january.csv", "sales_february.csv", "sales_march.csv"],
    )


class TestFormatMoney:
    @pytest.mark.parametrize(
        "value, expected",
        [
            (0, "0,00\u00a0₽"),
            (8126.5, "8\u00a0126,50\u00a0₽"),
            (-1500.0, "-1\u00a0500,00\u00a0₽"),
            (1234567.891, "1\u00a0234\u00a0567,89\u00a0₽"),
        ],
    )
    def test_formats_russian_style(self, value: float, expected: str) -> None:
        assert format_money(value) == expected


class TestPrintSummary:
    def test_prints_totals_and_dropped_rows(self, report: Report, capsys) -> None:
        print_summary(report)
        out = capsys.readouterr().out
        assert "Файлов: 3" in out
        assert "в отчёте: 25" in out
        assert "Отброшено 5 строк" in out
        assert "ИТОГО" in out
        assert "Кофе в зёрнах" in out

    def test_handles_empty_report(self, capsys) -> None:
        empty = pd.DataFrame(columns=list(CANONICAL_COLUMNS))
        from clean import CleanStats

        summary = summarize(empty, "Товар")
        print_summary(Report(summary, empty, CleanStats(), []))
        out = capsys.readouterr().out
        assert "ИТОГО" in out
        assert "0,00" in out


class TestSaveExcel:
    def test_writes_two_sheets(self, report: Report, tmp_path: Path) -> None:
        path = save_excel(report, tmp_path / "report.xlsx")
        assert path.exists() and path.stat().st_size > 0

        wb = load_workbook(path)
        assert wb.sheetnames == ["Сводка", "Данные"]

    def test_summary_sheet_has_title_and_totals(self, report: Report, tmp_path: Path) -> None:
        wb = load_workbook(save_excel(report, tmp_path / "report.xlsx"))
        ws = wb["Сводка"]

        assert "Отчёт по продажам" in ws["A1"].value
        values = [cell.value for row in ws.iter_rows(max_col=1) for cell in row]
        assert "ИТОГО" in values
        assert any(v == "Отброшено строк" for v in values)

    def test_data_sheet_has_real_dates_and_numbers(
        self, report: Report, tmp_path: Path
    ) -> None:
        wb = load_workbook(save_excel(report, tmp_path / "report.xlsx"))
        ws = wb["Данные"]

        assert ws["A2"].number_format == "DD.MM.YYYY"
        assert ws["A2"].value.year == 2025
        assert isinstance(ws["E2"].value, int)
        assert isinstance(ws["F2"].value, (int, float))
        assert ws.auto_filter.ref is not None

    def test_money_format_is_valid_excel(self, report: Report, tmp_path: Path) -> None:
        wb = load_workbook(save_excel(report, tmp_path / "report.xlsx"))
        ws = wb["Сводка"]

        money_cells = [
            cell
            for row in ws.iter_rows(min_row=20, max_col=6)
            for cell in row
            if cell.number_format != "General"
        ]
        assert money_cells, "форматы должны быть проставлены"

        # Деньги — с рублём, доля — с процентом; ни один формат не «голый» числовой.
        for cell in money_cells:
            fmt = cell.number_format
            assert "₽" in fmt or "%" in fmt, f"{cell.coordinate}: {fmt}"

    def test_share_column_shows_percent_sign(self, report: Report, tmp_path: Path) -> None:
        wb = load_workbook(save_excel(report, tmp_path / "report.xlsx"))
        ws = wb["Сводка"]

        header_row = next(
            row[0].row for row in ws.iter_rows(max_col=1) if row[0].value == "Группа"
        )
        share_cell = ws.cell(row=header_row + 1, column=6)
        assert share_cell.value == pytest.approx(50.73, abs=0.01)
        assert "%" in share_cell.number_format

    def test_creates_missing_directory(self, report: Report, tmp_path: Path) -> None:
        path = save_excel(report, tmp_path / "nested" / "dir" / "report.xlsx")
        assert path.exists()

    def test_empty_data_writes_file(self, tmp_path: Path) -> None:
        empty = pd.DataFrame(columns=list(CANONICAL_COLUMNS))
        from clean import CleanStats

        report = Report(summarize(empty, "product"), empty, CleanStats(), [])
        path = save_excel(report, tmp_path / "empty.xlsx")
        wb = load_workbook(path)
        assert wb.sheetnames == ["Сводка", "Данные"]
