"""Вывод результата: консольная сводка и Excel-файл с двумя листами."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from clean import CANONICAL_COLUMNS, CleanStats
from report import Summary

SUMMARY_COLUMNS = {
    "group": "Группа",
    "orders": "Заказов",
    "quantity": "Штук",
    "revenue": "Выручка, ₽",
    "avg_order_value": "Средний чек, ₽",
    "share_percent": "Доля, %",
}

DATA_COLUMNS = {
    "date": "Дата",
    "manager": "Менеджер",
    "product": "Товар",
    "city": "Город",
    "quantity": "Количество",
    "price": "Цена, ₽",
}


@dataclass
class Report:
    """Всё, что нужно для записи файла — собирается в одном месте."""

    summary: Summary
    data: pd.DataFrame
    stats: CleanStats
    source_names: list[str]


def format_money(value: float) -> str:
    """81 326,50 ₽ — с неразрывным пробелом между тысячами."""
    return f"{value:,.2f}".replace(",", "\u00a0").replace(".", ",") + "\u00a0₽"


def format_report_date(value: date | pd.Timestamp) -> str:
    return pd.Timestamp(value).strftime("%d.%m.%Y")


def _column_widths(ws: Worksheet, minimum: int = 12) -> None:
    """Ширина по содержимому, но не уже минимума."""
    for column_cells in ws.columns:
        length = max(
            (len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells
        )
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(
            minimum, min(length + 2, 60)
        )


def _write_summary_sheet(ws: Worksheet, report: Report) -> None:
    summary, totals, stats = report.summary, report.summary.totals, report.stats

    ws["A1"] = f"Отчёт по продажам — группировка: {summary.label}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Сформирован: {format_report_date(date.today())}"

    facts = [
        ("Файлов обработано", stats.files),
        ("Строк на входе", stats.input_rows),
        ("Строк в отчёте", stats.output_rows),
        ("Заказов", totals.orders),
        ("Продано, шт.", totals.quantity),
        ("Выручка", format_money(totals.revenue)),
        ("Средний чек", format_money(totals.avg_order_value)),
        ("Групп", totals.groups),
    ]

    row = 4
    for label, value in facts:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    dropped = stats.dropped_rows()
    if dropped:
        row += 1
        ws.cell(row=row, column=1, value="Отброшено строк").font = Font(bold=True)
        ws.cell(row=row, column=2, value=dropped)
        row += 1
        reasons = [
            ("пустые строки", stats.blank_rows),
            ("нераспознанная дата", stats.bad_date),
            ("нет или некорректное количество", stats.bad_quantity),
            ("нет или некорректная цена", stats.bad_price),
            ("полные дубли", stats.duplicates),
        ]
        for label, count in reasons:
            if count:
                ws.cell(row=row, column=1, value=f"  — {label}")
                ws.cell(row=row, column=2, value=count)
                row += 1

    if summary.warnings:
        row += 1
        for warning in summary.warnings:
            ws.cell(row=row, column=1, value=warning)
            row += 1

    start = row + 2
    for offset, name in enumerate(SUMMARY_COLUMNS.values()):
        cell = ws.cell(row=start, column=offset + 1, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F6F52")
        cell.alignment = Alignment(horizontal="center")

    for i, record in enumerate(summary.table.itertuples(index=False), start=start + 1):
        ws.cell(row=i, column=1, value=record.group)
        ws.cell(row=i, column=2, value=int(record.orders))
        ws.cell(row=i, column=3, value=int(record.quantity))
        ws.cell(row=i, column=4, value=float(record.revenue)).number_format = '#,##0.00 "₽"'
        ws.cell(row=i, column=5, value=float(record.avg_order_value)).number_format = (
            '#,##0.00 "₽"'
        )
        ws.cell(row=i, column=6, value=float(record.share_percent)).number_format = (
            '0.00"%"'
        )

    total_row = start + len(summary.table) + 1
    ws.cell(row=total_row, column=1, value="ИТОГО").font = Font(bold=True)
    ws.cell(row=total_row, column=2, value=totals.orders).font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=totals.quantity).font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=4, value=totals.revenue)
    total_cell.number_format = '#,##0.00 "₽"'
    total_cell.font = Font(bold=True)

    # Таблица ниже шапки: замораживать её бессмысленно, а шапку с фактами —
    # значит оставить на экране 20 строк и закрыть саму таблицу.
    _column_widths(ws)


def _write_data_sheet(ws: Worksheet, report: Report) -> None:
    headers = [DATA_COLUMNS[name] for name in CANONICAL_COLUMNS]
    for offset, name in enumerate(headers):
        cell = ws.cell(row=1, column=offset + 1, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F6F52")
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(report.data.itertuples(index=False), start=2):
        ws.cell(row=row_index, column=1, value=row.date.to_pydatetime()
                if hasattr(row.date, "to_pydatetime") else row.date).number_format = (
            "DD.MM.YYYY"
        )
        ws.cell(row=row_index, column=2, value=row.manager)
        ws.cell(row=row_index, column=3, value=row.product)
        ws.cell(row=row_index, column=4, value=row.city)
        ws.cell(row=row_index, column=5, value=int(row.quantity))
        ws.cell(row=row_index, column=6, value=float(row.price)).number_format = (
            '#,##0.00 "₽"'
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{max(len(report.data) + 1, 1)}"
    _column_widths(ws)


def save_excel(report: Report, path: str | Path) -> Path:
    """Пишет отчёт в .xlsx. Первый лист — сводка, второй — очищенные данные."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        writer.book.create_sheet("Сводка")
        writer.book.create_sheet("Данные")
        _write_summary_sheet(writer.book["Сводка"], report)
        _write_data_sheet(writer.book["Данные"], report)
        # openpyxl создаёт лист по умолчанию — убираем, чтобы файл начинался со «Сводки».
        if "Sheet" in writer.book.sheetnames:
            del writer.book["Sheet"]

    return target


def print_summary(report: Report) -> None:
    """Печатает сводку в консоль — для быстрой проверки без открытия Excel."""
    summary, totals, stats = report.summary, report.summary.totals, report.stats

    print()
    print(f"Файлов: {stats.files} | строк на входе: {stats.input_rows} | "
          f"в отчёте: {stats.output_rows}")

    if stats.encodings:
        encodings = ", ".join(f"{name}: {enc}" for name, enc in stats.encodings.items())
        print(f"Кодировки: {encodings}")

    if stats.ignored_columns:
        print(f"Лишние колонки пропущены: {', '.join(stats.ignored_columns)}")

    if stats.dropped_rows():
        reasons = [
            ("пустых", stats.blank_rows),
            ("без даты", stats.bad_date),
            ("без количества", stats.bad_quantity),
            ("без цены", stats.bad_price),
            ("дублей", stats.duplicates),
        ]
        detail = ", ".join(f"{label}: {count}" for label, count in reasons if count)
        print(f"Отброшено {stats.dropped_rows()} строк ({detail})")

    for warning in summary.warnings:
        print(f"Внимание: {warning}")

    print()
    header = (
        f"{'Группа':<24} {'Заказов':>8} {'Штук':>7} "
        f"{'Выручка':>15} {'Доля':>7}"
    )
    print(header)
    print("-" * len(header))

    for record in summary.table.itertuples(index=False):
        group = record.group if len(record.group) <= 23 else record.group[:20] + "..."
        print(
            f"{group:<24} {record.orders:>8} {record.quantity:>7} "
            f"{format_money(record.revenue):>15} {record.share_percent:>6.2f}%"
        )

    print("-" * len(header))
    print(
        f"{'ИТОГО':<24} {totals.orders:>8} {totals.quantity:>7} "
        f"{format_money(totals.revenue):>15}"
    )
    print(f"Средний чек: {format_money(totals.avg_order_value)}")
    print()
